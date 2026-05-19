"""Phase 2 烟测：赛道 + 条目 + Excel 导入导出全流程。"""
import io
import json
import sys
import urllib.parse
import requests
from openpyxl import Workbook, load_workbook

BASE = "http://127.0.0.1:8000"


def check(cond, msg):
    if not cond:
        print(f"  ✗ {msg}")
        sys.exit(1)
    print(f"  ✓ {msg}")


def main():
    # 0. 清干净（删历史 'Phase2 测试赛道'）
    r = requests.get(f"{BASE}/api/tracks")
    for t in r.json():
        if t["name"].startswith("Phase2"):
            requests.delete(f"{BASE}/api/tracks/{t['id']}")

    # 1. 创建赛道
    print("[1] 创建赛道")
    r = requests.post(f"{BASE}/api/tracks", json={"name": "Phase2 测试赛道", "description": "烟测"})
    check(r.status_code == 201, f"POST /api/tracks → 201 (got {r.status_code}: {r.text[:200]})")
    track = r.json()
    tid = track["id"]
    print(f"    id={tid[:8]}…  name={track['name']}  entry_count={track['entry_count']}")

    # 2. 重名应该 400
    r = requests.post(f"{BASE}/api/tracks", json={"name": "Phase2 测试赛道"})
    check(r.status_code == 400, "重名赛道 → 400")

    # 3. 列表
    r = requests.get(f"{BASE}/api/tracks")
    check(any(t["id"] == tid for t in r.json()), "GET /api/tracks 含新建赛道")

    # 4. 创建条目
    print("[2] 创建条目")
    r = requests.post(f"{BASE}/api/tracks/{tid}/entries", json={
        "title": "第一篇：电话恐惧症",
        "topic_direction": "代际差异",
        "publish_date": "2026-06-01",
        "status": "to_edit",
        "notes": "热点切入",
    })
    check(r.status_code == 201, f"POST entries → 201 (got {r.status_code}: {r.text[:200]})")
    e1 = r.json()
    print(f"    id={e1['id'][:8]}…  title={e1['title']}  pub={e1['publish_date']}")

    r = requests.post(f"{BASE}/api/tracks/{tid}/entries", json={"title": "第二篇：AI 焦虑"})
    e2 = r.json()
    r = requests.post(f"{BASE}/api/tracks/{tid}/entries", json={"title": "第三篇：Z 世代社交"})
    e3 = r.json()

    # 5. 列表 + 状态过滤
    r = requests.get(f"{BASE}/api/tracks/{tid}/entries")
    check(len(r.json()) == 3, f"GET entries → 3 (got {len(r.json())})")
    r = requests.get(f"{BASE}/api/tracks/{tid}/entries?status=to_edit")
    check(len(r.json()) == 3, "状态过滤 to_edit → 3")

    # 6. 单条更新
    print("[3] 更新单条")
    r = requests.patch(f"{BASE}/api/entries/{e1['id']}", json={"status": "to_publish", "notes": "已改稿"})
    check(r.status_code == 200 and r.json()["status"] == "to_publish", "PATCH 改状态 → to_publish")

    # 7. 批量改
    print("[4] 批量改状态")
    r = requests.patch(f"{BASE}/api/entries/batch", json={
        "ids": [e2["id"], e3["id"]],
        "status": "published",
    })
    check(r.json().get("updated") == 2, f"批量更新 → updated=2 (got {r.json()})")

    # 8. 删除单条
    print("[5] 删除单条")
    r = requests.delete(f"{BASE}/api/entries/{e3['id']}")
    check(r.status_code == 204, "DELETE entry → 204")
    r = requests.get(f"{BASE}/api/tracks/{tid}/entries")
    check(len(r.json()) == 2, f"剩余条目 → 2 (got {len(r.json())})")

    # 9. 导出 Excel
    print("[6] 导出 Excel")
    r = requests.get(f"{BASE}/api/tracks/{tid}/export")
    check(r.status_code == 200, f"GET export → 200 (got {r.status_code})")
    check("spreadsheetml" in r.headers.get("content-type", ""), "Content-Type 为 xlsx")
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    check(list(header) == ["标题", "选题方向", "发布日期", "状态", "备注"], f"表头正确 (got {header})")
    check(len(rows) - 1 == 2, f"数据行 = 2 (got {len(rows)-1})")
    print(f"    导出 {len(rows)-1} 行；表头={list(header)}")
    for r_ in rows[1:]:
        print(f"    {r_}")

    # 10. 导入 Excel
    print("[7] 导入 Excel")
    wb2 = Workbook()
    ws2 = wb2.active
    ws2.append(["标题", "选题方向", "发布日期", "状态", "备注"])
    ws2.append(["导入条目A", "方向A", "2026-07-01", "待编辑", "ok"])
    ws2.append(["导入条目B", "方向B", "2026-07-02", "已发布", ""])
    ws2.append(["", "空标题应跳过", "2026-07-03", "待编辑", ""])
    ws2.append(["导入条目C", None, None, "to_publish", "英文状态也接受"])
    buf = io.BytesIO()
    wb2.save(buf)
    buf.seek(0)

    r = requests.post(
        f"{BASE}/api/tracks/{tid}/import",
        files={"file": ("test.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    check(r.status_code == 200, f"POST import → 200 (got {r.status_code}: {r.text[:200]})")
    data = r.json()
    print(f"    {data}")
    check(data["created"] == 3, f"created=3 (got {data['created']})")
    check(data["skipped"] == 1, f"skipped=1 (got {data['skipped']})")

    # 11. 删除赛道（级联删条目）
    print("[8] 删除赛道（级联）")
    r = requests.delete(f"{BASE}/api/tracks/{tid}")
    check(r.status_code == 204, "DELETE track → 204")
    r = requests.get(f"{BASE}/api/tracks/{tid}/entries")
    check(r.status_code == 404, "条目列表 → 404（赛道已删）")

    print()
    print("✅ Phase 2 后端全流程通过")


if __name__ == "__main__":
    main()
