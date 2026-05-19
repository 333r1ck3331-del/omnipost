"""Phase 2 内容库路由测试 — 赛道 CRUD + 条目 CRUD + Excel 导入导出 + 路由顺序回归。"""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook, load_workbook


pytestmark = pytest.mark.asyncio


# ── Tracks ────────────────────────────────────────────────────────────


async def test_track_crud(client):
    # empty list
    r = await client.get("/api/tracks")
    assert r.status_code == 200
    assert isinstance(r.json(), list)

    # create
    r = await client.post("/api/tracks", json={"name": "T1", "description": "测试"})
    assert r.status_code == 201
    t = r.json()
    assert t["name"] == "T1" and t["entry_count"] == 0
    tid = t["id"]

    # duplicate name → 400
    r = await client.post("/api/tracks", json={"name": "T1"})
    assert r.status_code == 400

    # patch
    r = await client.patch(f"/api/tracks/{tid}", json={"description": "改了"})
    assert r.status_code == 200 and r.json()["description"] == "改了"

    # delete
    r = await client.delete(f"/api/tracks/{tid}")
    assert r.status_code == 204

    # 404
    r = await client.get(f"/api/tracks/{tid}/entries")
    assert r.status_code == 404


async def test_track_404_on_missing_id(client):
    r = await client.patch("/api/tracks/no-such-id", json={"name": "X"})
    assert r.status_code == 404
    r = await client.delete("/api/tracks/no-such-id")
    assert r.status_code == 404


# ── Entries ────────────────────────────────────────────────────────────


async def _new_track(client, name="TrackForEntries"):
    r = await client.post("/api/tracks", json={"name": name})
    return r.json()["id"]


async def test_entry_crud_and_filter(client):
    tid = await _new_track(client, "T-entries-1")

    # create 3
    r1 = (await client.post(f"/api/tracks/{tid}/entries",
                            json={"title": "A", "topic_direction": "方向A",
                                  "publish_date": "2026-06-01"})).json()
    r2 = (await client.post(f"/api/tracks/{tid}/entries", json={"title": "B"})).json()
    r3 = (await client.post(f"/api/tracks/{tid}/entries",
                            json={"title": "C", "status": "published"})).json()

    # list
    r = await client.get(f"/api/tracks/{tid}/entries")
    assert len(r.json()) == 3

    # filter
    r = await client.get(f"/api/tracks/{tid}/entries?status=to_edit")
    assert len(r.json()) == 2  # A and B default to_edit
    r = await client.get(f"/api/tracks/{tid}/entries?status=published")
    assert len(r.json()) == 1

    # patch single
    r = await client.patch(f"/api/entries/{r1['id']}",
                           json={"status": "to_publish", "notes": "已改"})
    assert r.status_code == 200
    assert r.json()["status"] == "to_publish" and r.json()["notes"] == "已改"

    # date is preserved as ISO string
    r = await client.get(f"/api/tracks/{tid}/entries")
    a = [e for e in r.json() if e["id"] == r1["id"]][0]
    assert a["publish_date"] == "2026-06-01"

    # delete
    r = await client.delete(f"/api/entries/{r2['id']}")
    assert r.status_code == 204
    r = await client.get(f"/api/tracks/{tid}/entries")
    assert len(r.json()) == 2

    # entry 404
    r = await client.patch("/api/entries/nope", json={"title": "x"})
    assert r.status_code == 404
    r = await client.delete("/api/entries/nope")
    assert r.status_code == 404


async def test_batch_update_route_order(client):
    """回归测试：/api/entries/batch 必须命中 batch 路由而不是 {entry_id}。

    这是 Phase 2 烟测时撞到的真实 bug — 修复后必须由测试守住。
    """
    tid = await _new_track(client, "T-batch")
    e1 = (await client.post(f"/api/tracks/{tid}/entries", json={"title": "X1"})).json()
    e2 = (await client.post(f"/api/tracks/{tid}/entries", json={"title": "X2"})).json()

    r = await client.patch("/api/entries/batch", json={
        "ids": [e1["id"], e2["id"]],
        "status": "published",
    })
    assert r.status_code == 200, f"batch 路由没命中（被 {{entry_id}} 拦了）: {r.text}"
    assert r.json()["updated"] == 2

    r = await client.get(f"/api/tracks/{tid}/entries?status=published")
    assert len(r.json()) == 2


async def test_cascade_delete(client):
    """删赛道时条目应级联删除。"""
    tid = await _new_track(client, "T-cascade")
    await client.post(f"/api/tracks/{tid}/entries", json={"title": "child1"})
    await client.post(f"/api/tracks/{tid}/entries", json={"title": "child2"})

    r = await client.delete(f"/api/tracks/{tid}")
    assert r.status_code == 204

    # entries also gone — query via API impossible (track 404)，直接查 DB
    from app.core.database import async_session
    from app.models import ContentEntry
    from sqlalchemy import select, func

    async with async_session() as s:
        n = (await s.execute(
            select(func.count()).select_from(ContentEntry).where(ContentEntry.track_id == tid)
        )).scalar_one()
    assert n == 0


# ── Excel 导出 / 导入 ──────────────────────────────────────────────────


async def test_export_xlsx(client):
    tid = await _new_track(client, "T-export")
    await client.post(f"/api/tracks/{tid}/entries", json={
        "title": "导出测试", "topic_direction": "方向X",
        "publish_date": "2026-07-01", "status": "to_publish", "notes": "备注",
    })
    await client.post(f"/api/tracks/{tid}/entries", json={"title": "第二条"})

    r = await client.get(f"/api/tracks/{tid}/export")
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]

    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert list(rows[0]) == ["标题", "选题方向", "发布日期", "状态", "备注"]
    assert len(rows) == 3
    # 第一行数据校验
    title, direction, pub, status, notes = rows[1]
    assert title == "导出测试"
    assert direction == "方向X"
    assert pub == "2026-07-01"
    assert status == "待发布"
    assert notes == "备注"


async def test_import_xlsx(client):
    tid = await _new_track(client, "T-import")

    wb = Workbook()
    ws = wb.active
    ws.append(["标题", "选题方向", "发布日期", "状态", "备注"])
    ws.append(["导入A", "方向A", "2026-08-01", "待编辑", "ok"])
    ws.append(["导入B", "", "2026-08-02", "已发布", ""])
    ws.append(["", "空标题跳过", "", "待编辑", ""])
    ws.append(["导入C", None, None, "to_publish", "英文状态"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    r = await client.post(
        f"/api/tracks/{tid}/import",
        files={"file": ("t.xlsx", buf.read(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["created"] == 3
    assert data["skipped"] == 1
    assert data["errors"] == []

    # 验证写入
    r = await client.get(f"/api/tracks/{tid}/entries")
    entries = r.json()
    assert len(entries) == 3
    titles = {e["title"] for e in entries}
    assert titles == {"导入A", "导入B", "导入C"}


async def test_import_rejects_non_xlsx(client):
    tid = await _new_track(client, "T-import-bad")
    r = await client.post(
        f"/api/tracks/{tid}/import",
        files={"file": ("x.csv", b"some,csv,data", "text/csv")},
    )
    assert r.status_code == 400
