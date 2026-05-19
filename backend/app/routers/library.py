"""Phase 2 — 内容库（赛道 + 选题清单）API.

- /api/tracks: 赛道 CRUD
- /api/tracks/{id}/entries: 赛道下的条目 CRUD
- /api/tracks/{id}/export: 导出 Excel (.xlsx)
- /api/tracks/{id}/import: 导入 Excel (.xlsx)
- /api/entries/{id}: 单条 CRUD
- /api/entries/batch: 批量改状态
"""

from __future__ import annotations

import io
import logging
from datetime import datetime, date
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy import func, select, update, delete
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.models import Track, ContentEntry
from app.schemas import (
    EntryBatchUpdate,
    EntryCreate,
    EntryOut,
    EntryUpdate,
    TrackCreate,
    TrackOut,
    TrackUpdate,
)

logger = logging.getLogger(__name__)
router = APIRouter(tags=["library"])


# ── Track helpers ───────────────────────────────────────────────────

async def _get_track(db: AsyncSession, track_id: str) -> Track:
    t = (await db.execute(select(Track).where(Track.id == track_id))).scalar_one_or_none()
    if not t:
        raise HTTPException(404, "赛道不存在")
    return t


async def _track_to_out(db: AsyncSession, t: Track) -> TrackOut:
    count = (await db.execute(
        select(func.count()).select_from(ContentEntry).where(ContentEntry.track_id == t.id)
    )).scalar_one()
    return TrackOut(
        id=t.id, name=t.name, description=t.description,
        sort_order=t.sort_order, entry_count=int(count), created_at=t.created_at,
    )


# ── Track routes ────────────────────────────────────────────────────

@router.get("/api/tracks", response_model=list[TrackOut])
async def list_tracks(db: AsyncSession = Depends(get_db)):
    rows = (await db.execute(
        select(Track).order_by(Track.sort_order, Track.created_at)
    )).scalars().all()
    return [await _track_to_out(db, t) for t in rows]


@router.post("/api/tracks", response_model=TrackOut, status_code=201)
async def create_track(body: TrackCreate, db: AsyncSession = Depends(get_db)):
    # check unique name
    existing = (await db.execute(select(Track).where(Track.name == body.name))).scalar_one_or_none()
    if existing:
        raise HTTPException(400, f"赛道名 '{body.name}' 已存在")
    # sort_order = current max + 1
    max_order = (await db.execute(select(func.coalesce(func.max(Track.sort_order), -1)))).scalar_one()
    t = Track(name=body.name, description=body.description, sort_order=int(max_order) + 1)
    db.add(t)
    await db.commit()
    await db.refresh(t)
    return await _track_to_out(db, t)


@router.patch("/api/tracks/{track_id}", response_model=TrackOut)
async def update_track(track_id: str, body: TrackUpdate, db: AsyncSession = Depends(get_db)):
    t = await _get_track(db, track_id)
    data = body.model_dump(exclude_unset=True)
    if "name" in data and data["name"] != t.name:
        dup = (await db.execute(select(Track).where(Track.name == data["name"]))).scalar_one_or_none()
        if dup:
            raise HTTPException(400, f"赛道名 '{data['name']}' 已存在")
    for k, v in data.items():
        setattr(t, k, v)
    await db.commit()
    await db.refresh(t)
    return await _track_to_out(db, t)


@router.delete("/api/tracks/{track_id}", status_code=204)
async def delete_track(track_id: str, db: AsyncSession = Depends(get_db)):
    t = await _get_track(db, track_id)
    await db.delete(t)
    await db.commit()
    return None


# ── Entry routes ────────────────────────────────────────────────────

@router.get("/api/tracks/{track_id}/entries", response_model=list[EntryOut])
async def list_entries(
    track_id: str,
    status: Optional[str] = None,
    db: AsyncSession = Depends(get_db),
):
    await _get_track(db, track_id)  # 404 if missing
    stmt = select(ContentEntry).where(ContentEntry.track_id == track_id)
    if status:
        stmt = stmt.where(ContentEntry.status == status)
    stmt = stmt.order_by(ContentEntry.sort_order, ContentEntry.created_at)
    rows = (await db.execute(stmt)).scalars().all()
    return list(rows)


@router.post("/api/tracks/{track_id}/entries", response_model=EntryOut, status_code=201)
async def create_entry(
    track_id: str,
    body: EntryCreate,
    db: AsyncSession = Depends(get_db),
):
    await _get_track(db, track_id)
    max_order = (await db.execute(
        select(func.coalesce(func.max(ContentEntry.sort_order), -1))
        .where(ContentEntry.track_id == track_id)
    )).scalar_one()
    entry = ContentEntry(
        track_id=track_id,
        title=body.title,
        topic_direction=body.topic_direction,
        publish_date=body.publish_date,
        status=body.status,
        notes=body.notes,
        sort_order=int(max_order) + 1,
    )
    db.add(entry)
    await db.commit()
    await db.refresh(entry)
    return entry


@router.patch("/api/entries/batch")
async def batch_update_entries(body: EntryBatchUpdate, db: AsyncSession = Depends(get_db)):
    res = await db.execute(
        update(ContentEntry)
        .where(ContentEntry.id.in_(body.ids))
        .values(status=body.status)
    )
    await db.commit()
    return {"updated": res.rowcount}


@router.patch("/api/entries/{entry_id}", response_model=EntryOut)
async def update_entry(entry_id: str, body: EntryUpdate, db: AsyncSession = Depends(get_db)):
    entry = (await db.execute(
        select(ContentEntry).where(ContentEntry.id == entry_id)
    )).scalar_one_or_none()
    if not entry:
        raise HTTPException(404, "条目不存在")
    data = body.model_dump(exclude_unset=True)
    for k, v in data.items():
        setattr(entry, k, v)
    await db.commit()
    await db.refresh(entry)
    return entry


@router.delete("/api/entries/{entry_id}", status_code=204)
async def delete_entry(entry_id: str, db: AsyncSession = Depends(get_db)):
    entry = (await db.execute(
        select(ContentEntry).where(ContentEntry.id == entry_id)
    )).scalar_one_or_none()
    if not entry:
        raise HTTPException(404, "条目不存在")
    await db.delete(entry)
    await db.commit()
    return None


# ── Excel import / export ───────────────────────────────────────────

_STATUS_CN = {"to_edit": "待编辑", "to_publish": "待发布", "published": "已发布"}
_STATUS_FROM_CN = {v: k for k, v in _STATUS_CN.items()}
# also accept English values directly
for k in list(_STATUS_CN.keys()):
    _STATUS_FROM_CN[k] = k


@router.get("/api/tracks/{track_id}/export")
async def export_track_xlsx(track_id: str, db: AsyncSession = Depends(get_db)):
    """导出该赛道下的全部条目为 .xlsx。"""
    from openpyxl import Workbook

    track = await _get_track(db, track_id)
    rows = (await db.execute(
        select(ContentEntry)
        .where(ContentEntry.track_id == track_id)
        .order_by(ContentEntry.sort_order, ContentEntry.created_at)
    )).scalars().all()

    wb = Workbook()
    ws = wb.active
    ws.title = (track.name or "track")[:31]
    ws.append(["标题", "选题方向", "发布日期", "状态", "备注"])
    for e in rows:
        ws.append([
            e.title,
            e.topic_direction or "",
            e.publish_date.isoformat() if e.publish_date else "",
            _STATUS_CN.get(e.status, e.status),
            e.notes or "",
        ])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    safe_name = track.name.replace("/", "_").replace("\\", "_")
    ts = datetime.now().strftime("%Y%m%d")
    filename = f"{safe_name}_{ts}.xlsx"
    # RFC 5987 — 中文文件名
    from urllib.parse import quote
    headers = {
        "Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"
    }
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


def _parse_date(v) -> Optional[date]:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        s = v.strip()
        if not s:
            return None
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                continue
    return None


@router.post("/api/tracks/{track_id}/import")
async def import_track_xlsx(
    track_id: str,
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
):
    """导入 .xlsx 到该赛道，第一行为表头：标题|选题方向|发布日期|状态|备注。

    标题为必填；空标题行会被跳过。状态接受中文（待编辑/待发布/已发布）或英文 key。
    """
    from openpyxl import load_workbook

    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(400, "仅支持 .xlsx 文件")
    await _get_track(db, track_id)

    raw = await file.read()
    if len(raw) > 5 * 1024 * 1024:
        raise HTTPException(400, "文件过大（>5MB）")

    try:
        wb = load_workbook(io.BytesIO(raw), data_only=True)
    except Exception as e:
        raise HTTPException(400, f"Excel 解析失败：{e}")
    ws = wb.active

    # current max sort_order
    max_order = (await db.execute(
        select(func.coalesce(func.max(ContentEntry.sort_order), -1))
        .where(ContentEntry.track_id == track_id)
    )).scalar_one()
    next_order = int(max_order) + 1

    created = 0
    skipped = 0
    errors: list[str] = []
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or all(v is None or v == "" for v in row):
            continue
        # pad to 5
        cells = list(row) + [None] * (5 - len(row))
        title, topic, pub, status, notes = cells[:5]
        if not title or not str(title).strip():
            skipped += 1
            continue
        status_key = _STATUS_FROM_CN.get(str(status).strip() if status else "", "to_edit")
        try:
            entry = ContentEntry(
                track_id=track_id,
                title=str(title).strip()[:200],
                topic_direction=str(topic).strip() if topic else None,
                publish_date=_parse_date(pub),
                status=status_key,
                notes=str(notes).strip() if notes else None,
                sort_order=next_order,
            )
            db.add(entry)
            next_order += 1
            created += 1
        except Exception as e:
            errors.append(f"第 {idx} 行：{e}")

    await db.commit()
    return {"created": created, "skipped": skipped, "errors": errors}
